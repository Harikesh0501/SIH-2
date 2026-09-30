import os
import sys
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import docx

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLE_DIR = os.path.join(ROOT_DIR, "sample_data")

def build_pdf(txt_filename, pdf_filename, doc_title, tag):
    txt_path = os.path.join(SAMPLE_DIR, txt_filename)
    pdf_path = os.path.join(SAMPLE_DIR, pdf_filename)
    
    if not os.path.exists(txt_path):
        print(f"[!] File not found: {txt_path}")
        return

    with open(txt_path, "r", encoding="utf-8") as f:
        lines = [line.strip() for line in f.readlines()]

    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#09090b'),
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#71717a'),
        spaceAfter=12
    )

    heading_style = ParagraphStyle(
        'DocHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#09090b'),
        spaceBefore=10,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#27272a'),
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#3f3f46'),
        leftIndent=15,
        spaceAfter=4
    )

    story = []

    # Header Tag
    story.append(Paragraph(f"<b>OFFICIAL DOCUMENT ARTIFACT</b> &bull; MoSPI / NSSTA Statistical Knowledge Repository &bull; {tag}", subtitle_style))
    story.append(Paragraph(doc_title, title_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#09090b'), spaceAfter=14))

    for line in lines:
        if not line:
            story.append(Spacer(1, 4))
            continue
        if line.startswith("GOVERNMENT OF INDIA") or line.startswith("MINISTRY OF") or line.startswith("NATIONAL STATISTICAL") or line.startswith("CENTRAL STATISTICS"):
            story.append(Paragraph(f"<b>{line}</b>", subtitle_style))
        elif line.startswith("CHAPTER") or line.startswith("MANUAL ON") or line.startswith("SOURCES AND") or line.startswith("PERIODIC LABOUR") or line.startswith("TRAINING HANDBOOK") or line.startswith("COMPILATION METHODOLOGY"):
            story.append(Paragraph(f"<b>{line}</b>", heading_style))
        elif len(line) > 2 and line[0].isdigit() and (line[1] == '.' or line[2] == '.'):
            story.append(Paragraph(f"<b>{line}</b>", heading_style))
        elif line.startswith("- ") or line.startswith("* "):
            story.append(Paragraph(f"&bull; {line[2:]}", bullet_style))
        elif line.startswith("a)") or line.startswith("b)") or line.startswith("c)") or line.startswith("1.") or line.startswith("2."):
            story.append(Paragraph(line, bullet_style))
        else:
            story.append(Paragraph(line, body_style))

    doc.build(story)
    print(f"[+] Compiled PDF: {pdf_filename}")

def build_docx(txt_filename, docx_filename, doc_title):
    txt_path = os.path.join(SAMPLE_DIR, txt_filename)
    docx_path = os.path.join(SAMPLE_DIR, docx_filename)
    
    if not os.path.exists(txt_path):
        print(f"[!] File not found: {txt_path}")
        return

    with open(txt_path, "r", encoding="utf-8") as f:
        lines = [line.strip() for line in f.readlines()]

    doc = docx.Document()
    
    # Title
    doc.add_heading(doc_title, level=0)
    doc.add_paragraph("Ministry of Statistics and Programme Implementation (MoSPI) | Official Reference Material")

    for line in lines:
        if not line:
            continue
        if len(line) > 2 and line[0].isdigit() and (line[1] == '.' or line[2] == '.'):
            doc.add_heading(line, level=1)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)

    doc.save(docx_path)
    print(f"[+] Compiled DOCX: {docx_filename}")

if __name__ == "__main__":
    print("[*] Generating MoSPI Sample Documents (PDF & DOCX)...")
    
    build_pdf(
        "cpi_manual_methodology_excerpt.txt",
        "cpi_manual_methodology_excerpt.pdf",
        "Manual on Consumer Price Index (CPI) - Compilation Methodology",
        "PRICE-STAT-001"
    )

    build_pdf(
        "national_accounts_gva_gdp_guide.txt",
        "national_accounts_gva_gdp_guide.pdf",
        "Sources & Methods: System of National Accounts (SNA 2008)",
        "MACRO-SNA-002"
    )

    build_docx(
        "plfs_sampling_and_concepts.txt",
        "plfs_sampling_and_concepts.docx",
        "Periodic Labour Force Survey (PLFS) - Survey Design & Estimation"
    )

    build_pdf(
        "plfs_sampling_and_concepts.txt",
        "plfs_sampling_and_concepts.pdf",
        "Periodic Labour Force Survey (PLFS) - Survey Design & Estimation",
        "SURVEY-PLFS-003"
    )

    build_pdf(
        "dqaf_and_statistical_ethics_guide.txt",
        "dqaf_and_statistical_ethics_guide.pdf",
        "Data Quality Assurance Framework (DQAF) & UN Fundamental Principles",
        "GOV-DQAF-004"
    )

    build_pdf(
        "iip_compilation_methodology_guide.txt",
        "iip_compilation_methodology_guide.pdf",
        "Index of Industrial Production (IIP) - Compilation Methodology",
        "IND-IIP-005"
    )

    print("[SUCCESS] All MoSPI sample documents packaged in /sample_data successfully!")
